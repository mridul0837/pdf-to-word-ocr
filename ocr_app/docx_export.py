"""Export (edited) preview text to an editable Word document."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Optional, Union

from docx import Document
from docx.shared import Inches, Pt

from .utils import is_page_marker, xml_safe


class ExportError(Exception):
    """The Word document could not be written; the message is user-facing."""


def split_pages(text: str) -> list[list[str]]:
    """Split preview text at page-marker lines; returns the lines of each page.

    Text before the first marker (or all text if there are no markers) forms
    its own page. Marker lines themselves are dropped.
    """
    pages: list[list[str]] = []
    current: Optional[list[str]] = None
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if is_page_marker(line):
            current = []
            pages.append(current)
            continue
        if current is None:
            current = []
            pages.append(current)
        current.append(line)
    # Drop a leading "page" that is only whitespace before the first marker.
    if len(pages) > 1 and not any(line.strip() for line in pages[0]):
        pages.pop(0)
    return pages


def build_document(text: str, title: Optional[str] = None) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(2)
    if title:
        doc.core_properties.title = xml_safe(title)

    for index, lines in enumerate(split_pages(text)):
        if index > 0:
            doc.add_page_break()
        # Trim leading/trailing blank lines; keep single blank lines as spacing.
        while lines and not lines[0].strip():
            lines.pop(0)
        while lines and not lines[-1].strip():
            lines.pop()
        previous_blank = False
        for line in lines:
            line = xml_safe(line.rstrip())
            if not line.strip():
                if not previous_blank:
                    doc.add_paragraph("")
                previous_blank = True
                continue
            doc.add_paragraph(line)
            previous_blank = False
    return doc


def export_docx(text: str, path: Union[str, os.PathLike], title: Optional[str] = None) -> Path:
    """Write ``text`` to a .docx file (page markers become page breaks).

    The file is written to a temporary name first and then moved into place,
    so a failed export never leaves a half-written document behind.
    """
    if not text or not text.strip():
        raise ExportError("There is no text to export.")
    path = Path(path)
    if path.suffix.lower() != ".docx":
        path = path.with_name(path.name + ".docx")
    if not path.parent.is_dir():
        raise ExportError(f"The folder does not exist: {path.parent}")

    doc = build_document(text, title)
    fd, tmp_name = tempfile.mkstemp(suffix=".docx", prefix=".~ocr_", dir=path.parent)
    os.close(fd)
    try:
        doc.save(tmp_name)
        os.replace(tmp_name, path)
    except PermissionError as exc:
        raise ExportError(
            f"Could not write {path.name}. If it is open in Word or another program, "
            f"close it and try again. ({exc})"
        ) from exc
    except OSError as exc:
        raise ExportError(f"Could not write {path}: {exc}") from exc
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)
    return path
