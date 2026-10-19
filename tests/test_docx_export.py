"""DOCX export: content, page breaks, sanitising and error handling."""

import os

import pytest
from docx import Document
from docx.oxml.ns import qn

from conftest import digital_pdf
from ocr_app import docx_export
from ocr_app.docx_export import ExportError, export_docx, split_pages
from ocr_app.pipeline import process_pdf


def paragraphs(path):
    return [p.text for p in Document(str(path)).paragraphs]


def page_break_count(path):
    body = Document(str(path)).element.body
    return sum(1 for br in body.iter(qn("w:br")) if br.get(qn("w:type")) == "page")


def test_export_with_page_markers(tmp_path):
    text = (
        "===== PAGE 1 · Embedded PDF text =====\nFirst line\nSecond line\n\n"
        "===== PAGE 2 · Tesseract OCR =====\nPage two text\n\n\n\nAfter gap\n\n"
        "===== PAGE 3 · TrOCR handwriting =====\nHandwritten words"
    )
    path = export_docx(text, tmp_path / "out.docx", title="My scan")
    assert path == tmp_path / "out.docx"
    content = [p for p in paragraphs(path) if p]
    assert content == ["First line", "Second line", "Page two text", "After gap", "Handwritten words"]
    assert not any("=====" in p for p in paragraphs(path))  # markers are not exported
    assert page_break_count(path) == 2
    assert Document(str(path)).core_properties.title == "My scan"


def test_blank_line_runs_collapse_to_one_empty_paragraph(tmp_path):
    path = export_docx("a\n\n\n\nb", tmp_path / "gap.docx")
    assert paragraphs(path) == ["a", "", "b"]


def test_text_without_markers_is_a_single_page(tmp_path):
    path = export_docx("Just some edited text\nwith two lines", tmp_path / "plain.docx")
    assert paragraphs(path) == ["Just some edited text", "with two lines"]
    assert page_break_count(path) == 0


def test_illegal_xml_characters_are_removed(tmp_path):
    path = export_docx("Bad\x00chars\x0b here\x0c\nnext\x1f line", tmp_path / "ctrl.docx")
    assert paragraphs(path) == ["Badchars here", "next line"]


def test_docx_extension_is_added(tmp_path):
    path = export_docx("hello", tmp_path / "report")
    assert path.name == "report.docx"
    assert path.is_file()


def test_empty_text_is_rejected(tmp_path):
    with pytest.raises(ExportError, match="no text"):
        export_docx("  \n ", tmp_path / "empty.docx")


def test_missing_folder_is_rejected(tmp_path):
    with pytest.raises(ExportError, match="folder does not exist"):
        export_docx("hello", tmp_path / "nope" / "out.docx")


def test_locked_file_gives_helpful_message_and_no_temp_files(tmp_path, monkeypatch):
    def locked(src, dst):
        raise PermissionError(13, "Permission denied", str(dst))

    monkeypatch.setattr(docx_export.os, "replace", locked)
    with pytest.raises(ExportError, match="open in Word"):
        export_docx("hello", tmp_path / "locked.docx")
    assert os.listdir(tmp_path) == []


def test_split_pages_handles_preamble_before_first_marker():
    pages = split_pages("intro\n===== PAGE 1 · x =====\none")
    assert pages == [["intro"], ["one"]]
    assert split_pages("\n===== PAGE 1 · x =====\none") == [["one"]]


def test_pipeline_to_docx_round_trip(tmp_path):
    texts = [f"Round trip page {i} has enough text to be embedded." for i in range(1, 4)]
    result = process_pdf(digital_pdf(texts))
    path = export_docx(result.to_text(), tmp_path / "round.docx")
    assert [p for p in paragraphs(path) if p] == texts
    assert page_break_count(path) == 2
